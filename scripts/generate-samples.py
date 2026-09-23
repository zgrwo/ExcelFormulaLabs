#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate-samples.py — 生成 samples/ExcelFormulaLabs-Samples.xlsx 示例工作簿。

用法：python scripts/generate-samples.py
依赖：openpyxl（仅生成时需要；生成物入库，用户无需安装）

说明：公式以字符串写入（openpyxl 不求值）。加载 .xll 后由 Excel 计算；
未加载时显示 #NAME? 属预期（见「使用说明」sheet）。
"""

import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("需要 openpyxl：pip install openpyxl", file=sys.stderr)
    sys.exit(1)

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "samples" / "ExcelFormulaLabs-Samples.xlsx"

TITLE_FONT = Font(bold=True, size=13)
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
LABEL_FONT = Font(bold=True)


def put_title(ws, text):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT


def put_table(ws, start_row, headers, rows):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
    for r, row in enumerate(rows, start=start_row + 1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    return start_row + len(rows)


def put_examples(ws, start_row, examples, label_col=6, formula_col=7):
    """examples: [(label, formula), ...]，写在第 label_col/formula_col 列。"""
    for i, (label, formula) in enumerate(examples):
        r = start_row + i
        lc = ws.cell(row=r, column=label_col, value=label)
        lc.font = LABEL_FONT
        ws.cell(row=r, column=formula_col, value=formula)
    return start_row + len(examples)


def autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build():
    wb = Workbook()

    # ── 使用说明 ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "使用说明"
    put_title(ws, "ExcelFormulaLabs 示例工作簿")
    notes = [
        "",
        "1. 先安装加载项：从 GitHub Release 下载对应位数的 .xll，运行 scripts/install.ps1 或按 README 手动加载。",
        "2. 本工作簿中所有示例公式在加载 .xll 后自动计算；未加载时显示 #NAME? 属预期。",
        "3. 32/64 位 Excel 必须匹配 .xll 位数；net48 版免安装运行时，net8.0 版需 .NET 8 Desktop Runtime。",
        "4. 各模块 sheet 左侧为示例数据，右侧（F/G 列）为公式与说明。",
        "5. 完整函数清单见 docs/specification/api-reference.md；逐函数示例见 docs/user-manual/user-manual.md。",
        "",
        "模块索引：" + "、".join(
            ["STATS", "STR", "REGEX", "DT", "ARR", "JSON", "XML", "DICT",
             "LINALG", "REGRESS", "SOLVE", "PHYCHEM", "DOE", "SQL", "PIVOT", "RANGE", "FS"]
        ),
    ]
    for i, t in enumerate(notes, start=2):
        ws.cell(row=i, column=1, value=t)
    autosize(ws, [110])

    # ── STATS ────────────────────────────────────────────────
    ws = wb.create_sheet("STATS")
    put_title(ws, "STATS.* — 统计（对标 scipy）")
    data = [12.5, 13.1, 11.8, 14.2, 13.7, 12.9, 13.5, 14.0, 12.2, 13.8]
    put_table(ws, 3, ["value"], [[v] for v in data])
    put_examples(ws, 3, [
        ("均值", "=STATS.MEAN(A4:A13)"),
        ("样本标准差", "=STATS.STDEV(A4:A13)"),
        ("中位数", "=STATS.MEDIAN(A4:A13)"),
        ("描述统计摘要（9 值溢出）", "=STATS.SUMMARY(A4:A13)"),
    ])
    autosize(ws, [12, 14, 14, 14, 4, 26, 26])

    # ── STR ──────────────────────────────────────────────────
    ws = wb.create_sheet("STR")
    put_title(ws, "STR.* — 字符串")
    put_table(ws, 3, ["text"], [["hello world"], ["Excel"], ["  spaced  "]])
    put_examples(ws, 3, [
        ("反转", "=STR.REVERSE(A4)"),
        ("首字母大写", "=STR.TITLE(A4)"),
        ("连接数组（忽略空）", '=STR.TEXTJOIN("-", TRUE, A4:A6)'),
    ])
    autosize(ws, [16, 14, 14, 4, 26, 30])

    # ── REGEX ────────────────────────────────────────────────
    ws = wb.create_sheet("REGEX")
    put_title(ws, "REGEX.* — 正则（Excel 原生没有）")
    put_table(ws, 3, ["text"], [["Order A-1024 x3"], ["no digits here"], ["tel 138-0000-0000"]])
    put_examples(ws, 3, [
        ("首个数字串", r'=REGEX.MATCH(A4, "\d+")'),
        ("是否含数字", r'=REGEX.ISMATCH(A4, "\d+")'),
        ("替换全部数字为 #", r'=REGEX.REPLACE(A4, "\d", "#")'),
    ])
    autosize(ws, [22, 14, 14, 4, 24, 32])

    # ── DT ───────────────────────────────────────────────────
    ws = wb.create_sheet("DT")
    put_title(ws, "DT.* — 日期时间")
    put_examples(ws, 3, [
        ("周岁（1990-05-15 → 今天）", "=DT.AGEYEARS(DATE(1990,5,15), TODAY())"),
        ("2024 是否闰年", "=DT.ISLEAP(2024)"),
        ("2026 复活节日期", "=DT.EASTER(2026)"),
    ])
    autosize(ws, [30, 4, 4, 4, 30, 34])

    # ── ARR ──────────────────────────────────────────────────
    ws = wb.create_sheet("ARR")
    put_title(ws, "ARR.* — 数组")
    put_table(ws, 3, ["value"], [[3], [1], [4], [1], [5], [9], [2]])
    put_examples(ws, 3, [
        ("去重（保留首次顺序）", "=ARR.UNIQUE(A4:A10)"),
        ("降序排列", "=ARR.SORTDESC(A4:A10)"),
        ("切片（从索引 1 起 3 个）", "=ARR.SLICE(A4:A10, 1, 3)"),
    ])
    autosize(ws, [12, 14, 14, 4, 26, 28])

    # ── JSON ─────────────────────────────────────────────────
    ws = wb.create_sheet("JSON")
    put_title(ws, "JSON.* — JSON 解析")
    put_table(ws, 3, ["json"], [['{"name":"Alice","scores":[90,85]}'], ["{bad json}"]])
    put_examples(ws, 3, [
        ("取字段", '=JSON.QUERY(A4, "name")'),
        ("取数组元素", '=JSON.QUERY(A4, "scores[0]")'),
        ("是否合法 JSON", "=JSON.VALIDATE(A4)"),
        ("美化输出", "=JSON.PRETTIFY(A4)"),
    ])
    autosize(ws, [36, 14, 14, 4, 22, 30])

    # ── XML ──────────────────────────────────────────────────
    ws = wb.create_sheet("XML")
    put_title(ws, "XML.* — XML / XPath")
    put_table(ws, 3, ["xml"], [["<r><n>1</n><n>2</n><n>3</n></r>"]])
    put_examples(ws, 3, [
        ("是否合法 XML", "=XML.VALIDATE(A4)"),
        ("XPath 查询（溢出）", '=XML.XPATH(A4, "//n")'),
    ])
    autosize(ws, [34, 14, 14, 4, 26, 30])

    # ── DICT ─────────────────────────────────────────────────
    ws = wb.create_sheet("DICT")
    put_title(ws, "DICT.* — 频率统计")
    put_table(ws, 3, ["item"], [["A"], ["B"], ["A"], ["C"], ["A"], ["B"]])
    put_examples(ws, 3, [
        ("频率统计（value/count 两列）", "=DICT.FREQUENCY(A4:A9)"),
    ])
    autosize(ws, [12, 14, 14, 4, 30, 30])

    # ── LINALG ───────────────────────────────────────────────
    ws = wb.create_sheet("LINALG")
    put_title(ws, "LINALG.* — 线性代数")
    put_table(ws, 3, ["c1", "c2", "c3"], [[4, 1, 2], [3, 5, 1], [2, 3, 6]])
    for i, v in enumerate([10, 12, 14], start=4):
        ws.cell(row=i, column=5, value=v)
    ws.cell(row=3, column=5, value="b").font = HEAD_FONT
    ws.cell(row=3, column=5).fill = HEAD_FILL
    put_examples(ws, 7, [
        ("行列式", "=LINALG.DET(A4:C6)"),
        ("解方程组 Ax=b", "=LINALG.SOLVE(A4:C6, E4:E6)"),
    ])
    autosize(ws, [8, 8, 8, 4, 8, 4, 30, 30])

    # ── REGRESS ──────────────────────────────────────────────
    ws = wb.create_sheet("REGRESS")
    put_title(ws, "REGRESS.* — 回归")
    put_table(ws, 3, ["y", "x1", "x2"], [
        [6.0, 1, 3], [6.0, 2, 1], [11.0, 3, 4], [11.0, 4, 2], [16.0, 5, 5],
    ])
    put_examples(ws, 3, [
        ("R²", "=REGRESS.RSQ(A4:A8, B4:C8)"),
        ("OLS 报告（11 行）", "=REGRESS.OLS(A4:A8, B4:C8)"),
    ])
    autosize(ws, [8, 8, 8, 4, 22, 30])

    # ── SOLVE ────────────────────────────────────────────────
    ws = wb.create_sheet("SOLVE")
    put_title(ws, "SOLVE.* — 工艺参数反解（历史 6 行 + 请求 1 行）")
    put_table(ws, 3, ["IncomingA", "VariableU1", "OutputY1"], [
        [5, 2, 7.5], [6, 3, 9.5], [7, 2, 8.5], [8, 4, 12.0], [9, 3, 11.0], [10, 2, 10.0],
        [10, None, 13.0],   # 请求行：u 留空待求，目标 13 → 解 u=4
    ])
    put_examples(ws, 3, [
        ("反解可调参数（请求行留空）", "=SOLVE.INVERSE(A4:C11)"),
    ])
    autosize(ws, [14, 14, 14, 4, 30, 30])

    # ── PHYCHEM ──────────────────────────────────────────────
    ws = wb.create_sheet("PHYCHEM")
    put_title(ws, "PHYCHEM.* — 物理化学换算")
    put_examples(ws, 3, [
        ("摄氏 → 华氏", "=PHYCHEM.C_TO_F(100)"),
        ("分子量 H2SO4", '=PHYCHEM.MOLWT("H2SO4")'),
        ("密度 = 质量/体积", "=PHYCHEM.DENSITY(100, 2)"),
    ])
    autosize(ws, [24, 4, 4, 4, 24, 28])

    # ── DOE ──────────────────────────────────────────────────
    ws = wb.create_sheet("DOE")
    put_title(ws, "DOE.* — 实验设计（全因子 2 因子 × 2 水平）")
    put_examples(ws, 3, [
        ("设计矩阵（含 StdOrder/RunOrder）", '=DOE.PLAN(2,2,0,2,"full",FALSE)'),
    ])
    autosize(ws, [32, 4, 4, 4, 34, 30])

    # ── SQL ──────────────────────────────────────────────────
    ws = wb.create_sheet("SQL")
    put_title(ws, "SQL.* — 对区域写 SQL")
    put_table(ws, 3, ["dept", "salary"], [
        ["Eng", 12000], ["Eng", 15000], ["Ops", 9000], ["Ops", 9500], ["QA", 11000],
    ])
    put_examples(ws, 3, [
        ("按部门求平均", '=SQL.QUERY(A4:B8, "SELECT dept, AVG(salary) FROM data GROUP BY dept")'),
    ])
    autosize(ws, [12, 12, 14, 4, 34, 56])

    # ── PIVOT ────────────────────────────────────────────────
    ws = wb.create_sheet("PIVOT")
    put_title(ws, "PIVOT.* — 分组聚合")
    put_table(ws, 3, ["region", "product", "amount"], [
        ["North", "A", 100], ["North", "B", 200], ["South", "A", 150],
        ["South", "B", 50], ["North", "A", 120],
    ])
    put_examples(ws, 3, [
        ("按 region 求和 amount", '=PIVOT.GROUPBY(A4:C8, {1}, 3, "sum")'),
    ])
    autosize(ws, [12, 12, 12, 4, 28, 40])

    # ── RANGE ────────────────────────────────────────────────
    ws = wb.create_sheet("RANGE")
    put_title(ws, "RANGE.* — 区域导出")
    put_table(ws, 3, ["name", "qty"], [["Widget", 3], ["Gadget", 7]])
    put_examples(ws, 3, [
        ("导出 Markdown", "=RANGE.TOMD(A4:B6, TRUE)"),
    ])
    autosize(ws, [12, 8, 14, 4, 22, 30])

    # ── FS ───────────────────────────────────────────────────
    ws = wb.create_sheet("FS")
    put_title(ws, "FS.* — 文件系统（纯路径函数示例，不读写文件）")
    put_examples(ws, 3, [
        ("取文件名（不含扩展名）", '=FS.BNAME("C:\\data\\report.xlsx")'),
        ("取扩展名", '=FS.EXT("C:\\data\\report.xlsx")'),
        ("规范化路径", '=FS.NORM("C:\\data\\..\\file.txt")'),
    ])
    autosize(ws, [28, 4, 4, 4, 30, 34])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"written: {OUT.relative_to(REPO)} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
